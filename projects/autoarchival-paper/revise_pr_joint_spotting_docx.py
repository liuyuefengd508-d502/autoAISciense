from __future__ import annotations

from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement


ROOT = Path(__file__).resolve().parent
DOC_PATH = ROOT / "pr_joint_submission_package" / "paper_draft_PR_joint_spotting.docx"


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def set_paragraph_text(paragraph, text: str, style: str | None = None):
    paragraph.clear()
    if style:
        paragraph.style = style
    paragraph.add_run(text)


def split_structured_paragraph(paragraph, items):
    """
    items: list[tuple[text, style]]
    Replace current paragraph with first item, then insert the rest after it.
    """
    first_text, first_style = items[0]
    set_paragraph_text(paragraph, first_text, first_style)
    current = paragraph
    for text, style in items[1:]:
        current = insert_paragraph_after(current, text, style)


def main():
    if not DOC_PATH.exists():
        raise FileNotFoundError(DOC_PATH)

    backup_path = DOC_PATH.with_suffix(".before_final_cleanup.bak.docx")
    shutil.copy2(DOC_PATH, backup_path)

    doc = Document(DOC_PATH)

    # 1) Targeted wording clean-up.
    replacements = {
        "Experimental results demonstrate that AutoArchival achieves a remarkable end-to-end Hmean of 75.3% on ICDAR 2015 and 76.8% on MCAD, The method exhibits significant advantages over baseline architectures on MCAD and maintains competitive performance on public benchmarks, all while requiring substantially fewer parameters (14.5M).":
            "Experimental results demonstrate that AutoArchival achieves an end-to-end Hmean of 75.3% on ICDAR 2015 and 76.8% on MCAD. The method shows clear advantages over baseline architectures on MCAD and maintains competitive performance on public benchmarks while requiring substantially fewer parameters (14.5M).",
        "As shown, AutoArchival achieves an E2E Hmean of 75.3% on IC15 and 76.1% on Total-Text, demonstrating highly competitive generalization. Notably, AutoArchival outperforms classical approaches like FOTS and Mask TextSpotter v3, and is competitive with ABCNet v2, while utilizing only 14.5M parameters (compared to ABCNet v2鈥檚 35M). This confirms that LLM-guided evolution can successfully distill redundant structural paths into a highly efficient network.":
            "As shown, AutoArchival achieves an E2E Hmean of 75.3% on IC15 and 76.1% on Total-Text, showing competitive generalization. Notably, AutoArchival outperforms classical approaches like FOTS and Mask TextSpotter v3, and remains competitive with ABCNet v2, while utilizing only 14.5M parameters (compared to ABCNet v2鈥檚 35M). This confirms that LLM-guided evolution can successfully distill redundant structural paths into a highly efficient network.",
        "We presented AutoArchival, a budget-constrained LLM-guided code-level architecture search framework designed to solve the challenges of historical document text spotting. By shifting the workload from empirical human design to AI evolution, the agent autonomously discovered highly specialized structural operators. Empirical results demonstrate that our autonomously synthesized network significantly outperforms traditional manual baselines on ICDAR 2015, Total-Text, and MCAD, achieving competitive accuracy with a strong efficiency-performance trade-off. This research establishes a promising paradigm for domain-specific visual architectures, suggesting the potential of LLM agents in navigating complex architectural search spaces.":
            "We presented AutoArchival, a budget-constrained LLM-guided code-level architecture search framework designed to address the challenges of historical document text spotting. By shifting the workload from empirical human design to AI-guided evolution, the agent discovered specialized structural operators for vertical archival text. Empirical results show that the searched network substantially outperforms the manually designed seed baseline and achieves competitive results on ICDAR 2015, Total-Text, and MCAD, with a strong efficiency-performance trade-off. Overall, this study highlights a useful direction for domain-specific visual architectures and suggests the potential of LLM agents for navigating complex architectural search spaces.",
    }
    for para in doc.paragraphs:
        text = para.text.strip()
        if text in replacements:
            set_paragraph_text(para, replacements[text], para.style.name if para.style else None)

    # 2) Split heading/body paragraphs into cleaner journal-style structure.
    for para in list(doc.paragraphs):
        text = para.text.strip()
        if text == "Abstract\n" or not text:
            continue

        if text.startswith("Abstract\n"):
            body = text.split("\n", 1)[1].strip()
            split_structured_paragraph(para, [("Abstract", "Heading 1"), (body, "Normal")])
            continue

        if text.startswith("1. Introduction\n"):
            body = text.split("\n", 1)[1].strip()
            split_structured_paragraph(para, [("1. Introduction", "Heading 1"), (body, "Normal")])
            continue

        if text.startswith("2. Related Work\n2.1 Historical Document Analysis "):
            body = text.replace("2. Related Work\n2.1 Historical Document Analysis ", "", 1).strip()
            split_structured_paragraph(
                para,
                [("2. Related Work", "Heading 1"), ("2.1 Historical Document Analysis", "Heading 2"), (body, "Normal")],
            )
            continue

        if text.startswith("2.2 Joint Text Detection and Recognition\n"):
            body = text.split("\n", 1)[1].strip()
            split_structured_paragraph(para, [("2.2 Joint Text Detection and Recognition", "Heading 2"), (body, "Normal")])
            continue

        if text.startswith("2.3 LLM Agents for Code Generation and Research\n"):
            body = text.split("\n", 1)[1].strip()
            split_structured_paragraph(para, [("2.3 LLM Agents for Code Generation and Research", "Heading 2"), (body, "Normal")])
            continue

        if text.startswith("3. Proposed Method\n3.1 Problem Definition\n"):
            body = text.replace("3. Proposed Method\n3.1 Problem Definition\n", "", 1).strip()
            split_structured_paragraph(
                para,
                [("3. Proposed Method", "Heading 1"), ("3.1 Problem Definition", "Heading 2"), (body, "Normal")],
            )
            continue

        for heading in [
            "3.2 Baseline Joint Architecture",
            "3.3 Search Space and Editable Components",
            "3.4 LLM-Guided Autonomous Search Loop",
            "3.5 Joint Optimization Objective and Fitness Function",
            "3.6 Characteristics of the Searched Architecture",
            "4.2 Implementation Details",
            "4.3 Evaluation Metrics",
            "5.2 Results on MCAD",
            "5.3 Ablation Study",
            "5.4 Efficiency Analysis",
            "5.5 Qualitative Results",
            "5.6 Failure Cases and Discussion",
            "6. Conclusion",
        ]:
            prefix = heading + "\n"
            if text.startswith(prefix):
                body = text[len(prefix):].strip()
                split_structured_paragraph(
                    para,
                    [(heading, "Heading 2" if not heading.startswith("6.") else "Heading 1"), (body, "Normal")],
                )
                break
        else:
            if text.startswith("4. Experimental Setup\n4.1 Datasets\n"):
                body = text.replace("4. Experimental Setup\n4.1 Datasets\n", "", 1).strip()
                split_structured_paragraph(
                    para,
                    [("4. Experimental Setup", "Heading 1"), ("4.1 Datasets", "Heading 2"), (body, "Normal")],
                )
                continue
            if text.startswith("5. Results and Discussion\n5.1 Comparison with State-of-the-Art Methods\n"):
                body = text.replace("5. Results and Discussion\n5.1 Comparison with State-of-the-Art Methods\n", "", 1).strip()
                split_structured_paragraph(
                    para,
                    [("5. Results and Discussion", "Heading 1"), ("5.1 Comparison with State-of-the-Art Methods", "Heading 2"), (body, "Normal")],
                )
                continue
            if text.startswith("References "):
                refs_text = text.replace("References ", "", 1)
                refs = [chunk.strip() for chunk in refs_text.split(" [") if chunk.strip()]
                normalized = []
                for idx, ref in enumerate(refs):
                    if idx == 0 and ref.startswith("[1]"):
                        normalized.append(ref)
                    else:
                        normalized.append("[" + ref if not ref.startswith("[") else ref)
                items = [("References", "Heading 1")] + [(ref, "Normal") for ref in normalized]
                split_structured_paragraph(para, items)
                continue

    doc.save(DOC_PATH)
    print(f"Updated {DOC_PATH}")
    print(f"Backup saved to {backup_path}")


if __name__ == "__main__":
    main()
